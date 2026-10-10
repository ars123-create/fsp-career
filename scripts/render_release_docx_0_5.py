"""Editable technical documentation; render with the bundled documents helper."""
from pathlib import Path
import argparse
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


def plain(text):
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def build(target):
    if target.exists():
        raise FileExistsError("Choose a new output; preserve the reviewed document")
    doc = Document()
    doc.core_properties.title = "ФСП Карьера Техническая документация"
    doc.core_properties.author = "Лобутев Арсений Константинович; Просветов Александр Сергеевич"
    doc.core_properties.last_modified_by = ""
    doc.core_properties.subject = "Прототип 0.5.0"
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.9)
    section.left_margin = section.right_margin = Cm(2.1)
    for name, size in [("Normal", 10.5), ("Title", 24), ("Subtitle", 13),
                       ("Heading 1", 15), ("Heading 2", 12), ("Header", 8), ("Footer", 8)]:
        style = doc.styles[name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "DejaVu Sans")
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.13
    doc.styles["Normal"].paragraph_format.widow_control = True
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(14)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    # Remove the Word template's paragraph rule instead of inheriting its accent.
    for border in list(doc.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    doc.add_paragraph("ФСП Карьера", "Title")
    doc.add_paragraph("Техническая документация", "Subtitle")
    lines = (ROOT / "docs/TECHNICAL_0_5.md").read_text().splitlines()[2:]
    paragraph = []

    def flush():
        if paragraph:
            doc.add_paragraph(plain(" ".join(paragraph)))
            paragraph.clear()

    for line in lines:
        if line.startswith("## "):
            flush()
            title = re.sub(r"^(\d+)\.\s*", r"\1 ", line[3:]).replace(" — ", " ")
            doc.add_paragraph(title, "Heading 1")
        elif not line.strip():
            flush()
        else:
            paragraph.append(line.strip())
    flush()
    header = section.header.paragraphs[0]
    header.text = "ФСП Карьера · 0.5.0 · 10 октября 2026"
    footer = section.footer.paragraphs[0]
    footer.text = "Демонстрационный прототип · "
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.save(target)
    reopened = Document(target)
    assert len(reopened.paragraphs) == len(doc.paragraphs)
    assert "0.5.0" in " ".join(p.text for p in reopened.paragraphs)
    return len(reopened.paragraphs)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    print({"path": str(args.output), "paragraphs": build(args.output)})
